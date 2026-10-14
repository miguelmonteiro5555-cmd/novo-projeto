"""Base de cálculo mensal de fundo público (modelo FEMTA).

Gera a planilha com uma aba por mês e o "Resumo Anual", no mesmo layout do
modelo usado pelo fundo: comparativo da receita (prevista x atualizada x
realizada), da despesa (dotação inicial x autorizada x execução), o balanço
financeiro (ingressos x dispêndios), o RFP e as conferências.

Uso:
    python -m balancete.base_calculo --de base_existente.xlsx --saida nova.xlsx
    python -m balancete.base_calculo --ano 2026 --entidade "Fundo X" --saida modelo.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]

# Campo digitado -> célula na aba do mês.
CAMPOS = {
    # Quadro 1 — Receita
    "receita_prevista": "C5", "receita_atualizada": "D5", "receita_realizada": "E5",
    "superavit_previsto": "C7", "superavit_atualizado": "D7", "superavit_realizado": "E7",
    "excesso_previsto": "C8", "excesso_atualizado": "D8", "excesso_realizado": "E8",
    "saldo_mes_ant_previsto": "C9", "saldo_mes_ant_atualizado": "D9", "saldo_mes_ant_realizado": "E9",
    # Quadro 2 — Despesa
    "correntes_inicial": "C16", "correntes_autorizada": "D16", "correntes_execucao": "E16",
    "capital_inicial": "C17", "capital_autorizada": "D17", "capital_execucao": "E17",
    "reserva_inicial": "C19", "reserva_autorizada": "D19", "reserva_execucao": "E19",
    "suplementacao_inicial": "C20", "suplementacao_autorizada": "D20", "suplementacao_execucao": "E20",
    "especial_inicial": "C21", "especial_autorizada": "D21", "especial_execucao": "E21",
    "anulacoes_inicial": "C22", "anulacoes_autorizada": "D22", "anulacoes_execucao": "E22",
    # Quadro 3 — Balanço financeiro
    "transf_recebida_orcamentaria": "K6", "receita_extraorcamentaria": "K7",
    "transf_recebida_extraorcamentaria": "K8", "saldo_exercicio_anterior": "K9",
    "transf_concedida_orcamentaria": "M6", "despesa_extraorcamentaria": "M7",
    "transf_concedida_extraorcamentaria": "M8", "saldo_final": "M9",
}

ZERO = Decimal(0)


class ErroBase(Exception):
    pass


# ---------------------------------------------------------------- leitura

_FORMULA_NUMERICA = re.compile(r"^=\s*[-+]?[\d.,]+(\s*[-+]\s*[\d.,]+)*\s*$")


def _numero(valor) -> Decimal:
    """Número da célula; aceita fórmulas simples como "=3745677.7+66760.56"."""
    if valor is None or valor == "":
        return ZERO
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = str(valor).strip()
    if _FORMULA_NUMERICA.match(texto):
        total = ZERO
        for sinal, numero in re.findall(r"([-+]?)\s*([\d.]+)", texto[1:]):
            total += -Decimal(numero) if sinal == "-" else Decimal(numero)
        return total
    if texto.startswith("="):
        raise ValueError(f"fórmula não suportada: {texto}")
    if not texto.strip():
        return ZERO
    return Decimal(texto.replace(".", "").replace(",", ".")) if "," in texto else Decimal(texto)


def ler_base(caminho: str | Path) -> dict:
    """Lê os valores digitados de uma base no layout do modelo (abas jan..dez)."""
    from openpyxl import load_workbook

    livro = load_workbook(caminho)
    nomes = {ws.title.strip().lower(): ws for ws in livro.worksheets}
    meses, problemas = {}, []
    for mes in MESES:
        aba = nomes.get(mes)
        if aba is None:
            meses[mes] = {campo: ZERO for campo in CAMPOS}
            continue
        valores = {}
        for campo, celula in CAMPOS.items():
            if mes != "jan" and campo == "saldo_exercicio_anterior":
                valores[campo] = ZERO  # vem do saldo final do mês anterior
                continue
            try:
                valores[campo] = _numero(aba[celula].value)
            except ValueError as erro:
                problemas.append(f"Aba {mes}, célula {celula}: {erro}")
        meses[mes] = valores
    if problemas:
        raise ErroBase("\n".join(problemas))

    resumo = nomes.get("resumo anual")
    ano = None
    if resumo is not None and isinstance(resumo["D4"].value, datetime):
        ano = resumo["D4"].value.year
    return {"ano": ano, "meses": meses}


# ---------------------------------------------------------------- cálculo

def calcular(meses: dict[str, dict]) -> dict[str, dict]:
    """Valores calculados de cada mês (os mesmos das fórmulas da planilha)."""
    resultado = {}
    anterior = None
    for mes in MESES:
        v = meses[mes]
        r: dict[str, Decimal] = {}
        r["despesa_total_execucao"] = v["correntes_execucao"] + v["capital_execucao"]
        r["dotacao_inicial_total"] = v["correntes_inicial"] + v["capital_inicial"]
        r["dotacao_autorizada_total"] = v["correntes_autorizada"] + v["capital_autorizada"]
        r["despesa_orcamentaria"] = r["despesa_total_execucao"] - v["anulacoes_execucao"]
        r["saldo_anterior"] = v["saldo_exercicio_anterior"] if anterior is None else resultado[anterior]["saldo_final"]
        r["saldo_final"] = v["saldo_final"]
        r["ingressos"] = (v["receita_realizada"] + v["transf_recebida_orcamentaria"] + v["receita_extraorcamentaria"]
                          + v["transf_recebida_extraorcamentaria"] + r["saldo_anterior"])
        r["dispendios"] = (r["despesa_orcamentaria"] + v["transf_concedida_orcamentaria"] + v["despesa_extraorcamentaria"]
                           + v["transf_concedida_extraorcamentaria"] + v["saldo_final"])
        r["ro_re"] = v["receita_realizada"] + v["receita_extraorcamentaria"]
        r["do_de"] = r["despesa_orcamentaria"] + v["despesa_extraorcamentaria"]
        r["rfp"] = r["ro_re"] - r["do_de"]
        r["diferenca_quadros"] = v["receita_realizada"] - r["despesa_total_execucao"]
        r["receita_acumulada"] = v["receita_realizada"] + (resultado[anterior]["receita_acumulada"] if anterior else ZERO)
        r["despesa_acumulada"] = r["despesa_orcamentaria"] + (resultado[anterior]["despesa_acumulada"] if anterior else ZERO)
        r["conferencia_balanco"] = r["ingressos"] - r["dispendios"]
        r["conferencia_dotacao"] = r["dotacao_autorizada_total"] - (
            r["dotacao_inicial_total"] + v["suplementacao_autorizada"] - v["anulacoes_autorizada"])
        resultado[mes] = r
        anterior = mes
    return resultado


def conferir(meses: dict[str, dict]) -> list[str]:
    """Avisos das conferências automáticas (só meses preenchidos)."""
    avisos = []
    calculado = calcular(meses)
    for mes in MESES:
        if not any(meses[mes].values()):
            continue
        r = calculado[mes]
        if abs(r["conferencia_balanco"]) >= Decimal("0.01"):
            avisos.append(f"{mes}: ingressos e dispêndios diferem em {r['conferencia_balanco']:,.2f}")
        if abs(r["conferencia_dotacao"]) >= Decimal("0.01"):
            avisos.append(f"{mes}: dotação autorizada difere de inicial + suplementações − anulações "
                          f"em {r['conferencia_dotacao']:,.2f}")
    return avisos


# ---------------------------------------------------------------- planilha

MOEDA = '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-'
MOEDA_SIMPLES = '"R$"\\ #,##0.00;[Red]\\-"R$"\\ #,##0.00'
PERCENTUAL = "0.00%"


def gerar_base(meses: dict[str, dict], caminho: str | Path, ano: int, entidade: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    fino = Side(style="thin")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    cinza_escuro = PatternFill("solid", fgColor="A6A6A6")
    cinza_claro = PatternFill("solid", fgColor="D9D9D9")
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    meio = Alignment(vertical="center", wrap_text=True)

    def estilo(aba, ref, fonte, *, negrito=False, fundo=None, formato=None, alinhamento=meio):
        celula = aba[ref]
        celula.font = Font(name=fonte, size=12 if fonte == "Arial" else 11, bold=negrito)
        celula.border = borda
        celula.alignment = alinhamento
        if fundo:
            celula.fill = fundo
        if formato:
            celula.number_format = formato
        return celula

    def escrever(aba, ref, valor, fonte="Arial", **kw):
        estilo(aba, ref, fonte, **kw).value = valor

    livro = Workbook()
    livro.remove(livro.active)

    for i, mes in enumerate(MESES):
        aba = livro.create_sheet(mes)
        aba.sheet_view.showGridLines = False
        aba.sheet_view.zoomScale = 84
        anterior = MESES[i - 1] if i else None
        v = meses[mes]

        larguras = {"A": 9, "B": 36.6, "C": 26.1, "D": 28.7, "E": 20.9, "F": 16, "G": 20.1,
                    "H": 9, "I": 4, "J": 37.4, "K": 22.4, "L": 41.3, "M": 22.4}
        for coluna, largura in larguras.items():
            aba.column_dimensions[coluna].width = largura

        if entidade:
            aba["B1"] = f"{entidade} — Base de cálculo {mes.upper()}/{ano}"
            aba["B1"].font = Font(name="Arial", size=12, bold=True)

        # Quadro 1 — Receita
        for ref, texto in {"B3": "TÍTULOS", "C3": "RECEITA PREVISTA", "D3": "RECEITA ATUALIZADA",
                           "E3": "REALIZAÇÃO", "F3": "% PREVISTO", "G3": "% ATUALIZADO"}.items():
            escrever(aba, ref, texto, negrito=True, fundo=cinza_escuro, alinhamento=centro)
            escrever(aba, ref[0] + "4", None, fundo=cinza_escuro)
            aba.merge_cells(f"{ref}:{ref[0]}4")
        linhas_receita = {5: "Receitas Correntes", 6: "Total", 7: "Superávit Fin. Exer. Anterior",
                          8: "Excesso de Arrecadação", 9: "Saldo do Mês Anterior", 10: "Total Acumulado"}
        for linha, titulo in linhas_receita.items():
            destaque = linha in (6, 10)
            fundo = cinza_claro if destaque else None
            escrever(aba, f"B{linha}", titulo, negrito=destaque, fundo=fundo)
            for col in "CDE":
                ref = f"{col}{linha}"
                if linha == 6:
                    valor = f"={col}5"
                elif linha == 10:
                    if col == "E":
                        valor = f"={anterior}!E10+{mes}!E6" if anterior else "=E6"
                    else:
                        valor = f"={col}6"
                else:
                    campo = next(c for c, r in CAMPOS.items() if r == ref)
                    valor = float(v[campo])
                escrever(aba, ref, valor, negrito=destaque, fundo=fundo, formato=MOEDA,
                         alinhamento=centro if linha == 10 else meio)
            escrever(aba, f"F{linha}", f"=IFERROR(E{linha}/C{linha},0)", negrito=destaque, fundo=fundo,
                     formato=PERCENTUAL, alinhamento=centro)
            escrever(aba, f"G{linha}", f"=IFERROR(E{linha}/D{linha},0)", negrito=destaque, fundo=fundo,
                     formato=PERCENTUAL, alinhamento=centro)

        # Quadro 3 — Balanço financeiro
        escrever(aba, "J3", "RECEITA", negrito=True, fundo=cinza_escuro, alinhamento=centro)
        escrever(aba, "K3", None, fundo=cinza_escuro)
        aba.merge_cells("J3:K3")
        escrever(aba, "L3", "DESPESA", negrito=True, fundo=cinza_escuro, alinhamento=centro)
        escrever(aba, "M3", None, fundo=cinza_escuro)
        aba.merge_cells("L3:M3")
        escrever(aba, "J4", "INGRESSOS", negrito=True, fundo=cinza_claro, alinhamento=centro)
        escrever(aba, "K4", "=SUM(K5:K9)", negrito=True, fundo=cinza_claro, formato=MOEDA_SIMPLES, alinhamento=centro)
        escrever(aba, "L4", "DISPÊNDIOS", negrito=True, fundo=cinza_claro, alinhamento=centro)
        escrever(aba, "M4", "=SUM(M5:M9)", negrito=True, fundo=cinza_claro, formato=MOEDA_SIMPLES, alinhamento=centro)
        rotulos_ingresso = {5: "Receita Orçamentária", 6: "Transf. Recebida Orçamentária",
                            7: "Receita Extraorçamentária", 8: "Transf. Recebida Extraorçamentária",
                            9: "Saldo do Exercício Anterior" if not anterior else "Saldo do Mês Anterior"}
        rotulos_dispendio = {5: "Despesa Orçamentária", 6: "Transf. Concedida Orçamentária",
                             7: "Despesa Extraorçamentária", 8: "Transf. Concedida Extraorçamentária",
                             9: "Saldo para o Exercício Seguinte"}
        for linha in range(5, 10):
            escrever(aba, f"J{linha}", rotulos_ingresso[linha])
            escrever(aba, f"L{linha}", rotulos_dispendio[linha])
            if linha == 5:
                k, m = "=E6", "=E18-E22"
            elif linha == 9 and anterior:
                k, m = f"={anterior}!M9", float(v["saldo_final"])
            else:
                k = float(v[next(c for c, r in CAMPOS.items() if r == f"K{linha}")])
                m = float(v[next(c for c, r in CAMPOS.items() if r == f"M{linha}")])
            escrever(aba, f"K{linha}", k, formato=MOEDA, alinhamento=centro)
            escrever(aba, f"M{linha}", m, formato=MOEDA, alinhamento=centro)

        # Quadro 2 — Despesa
        for ref, texto in {"B14": "TÍTULOS", "C14": "DOTAÇÃO INICIAL", "D14": "DOTAÇÃO AUTORIZADA",
                           "E14": "EXECUÇÃO", "F14": "PERCENTUAL"}.items():
            escrever(aba, ref, texto, negrito=True, fundo=cinza_escuro, alinhamento=centro)
            escrever(aba, ref[0] + "15", None, fundo=cinza_escuro)
            aba.merge_cells(f"{ref}:{ref[0]}15")
        linhas_despesa = {16: "Despesas Correntes", 17: "Despesas de Capital e Amortização", 18: "Total",
                          19: "Reserva de Contingência", 20: "Suplementação", 21: "Dotação Especial",
                          22: "Anulações", 23: "Total Acumulado"}
        for linha, titulo in linhas_despesa.items():
            destaque = linha in (18, 23)
            fundo = cinza_claro if destaque else None
            escrever(aba, f"B{linha}", titulo, negrito=destaque, fundo=fundo)
            for col in "CDE":
                ref = f"{col}{linha}"
                if linha == 18:
                    valor = f"={col}16+{col}17"
                elif linha == 23:
                    valor = {
                        "C": "=C18+C19+C20+C21-C22",
                        # A dotação autorizada já inclui suplementações e anulações.
                        "D": "=D18+D19",
                        "E": f"={anterior}!E23+{mes}!E18-{mes}!E22" if anterior else "=E18-E22",
                    }[col]
                else:
                    valor = float(v[next(c for c, r in CAMPOS.items() if r == ref)])
                escrever(aba, ref, valor, negrito=destaque, fundo=fundo, formato=MOEDA)
            formula_f = "=IFERROR(E23/C23,0)" if linha == 23 else f"=IFERROR(E{linha}/D{linha},0)"
            escrever(aba, f"F{linha}", formula_f, negrito=destaque, fundo=fundo, formato=PERCENTUAL, alinhamento=centro)

        # RFP e conferências
        fundo_rfp = PatternFill("solid", fgColor="BFBFBF")
        escrever(aba, "J14", "RFP = (RO + RE) - (DO + DE)", negrito=True, fundo=fundo_rfp, alinhamento=centro)
        escrever(aba, "K14", None, fundo=fundo_rfp)
        aba.merge_cells("J14:K14")
        for linha, (formula, rotulo, formato) in {
            15: ("=K5+K7", "(RO + RE)", MOEDA),
            16: ("=M5+M7", "(DO + DE)", MOEDA),
            17: ("=J15-J16", "RFP", MOEDA_SIMPLES),
        }.items():
            escrever(aba, f"J{linha}", formula, negrito=linha == 17, formato=formato, alinhamento=centro)
            escrever(aba, f"K{linha}", rotulo, negrito=True, alinhamento=centro)
        escrever(aba, "J20", "DIFERENÇA DO QUADRO 1 E 2", negrito=True, fundo=cinza_claro, alinhamento=centro)
        escrever(aba, "J21", "=E6-E18", negrito=True, formato=MOEDA, alinhamento=centro)

        escrever(aba, "L14", "CONFERÊNCIAS", negrito=True, fundo=fundo_rfp, alinhamento=centro)
        escrever(aba, "M14", None, fundo=fundo_rfp)
        aba.merge_cells("L14:M14")
        escrever(aba, "L15", "Ingressos − Dispêndios")
        escrever(aba, "M15", "=ROUND(K4-M4,2)", formato=MOEDA, alinhamento=centro)
        escrever(aba, "L16", "Dotação autorizada − (inicial + supl. − anul.)")
        escrever(aba, "M16", "=ROUND(D18-(C18+D20-D22),2)", formato=MOEDA, alinhamento=centro)
        escrever(aba, "L17", "Situação", negrito=True)
        escrever(aba, "M17", '=IF(AND(M15=0,M16=0),"OK","VERIFICAR")', negrito=True, alinhamento=centro)

    _resumo_anual(livro, ano, escrever, cinza_escuro, cinza_claro, centro)

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(caminho)
    return caminho


def _resumo_anual(livro, ano, escrever, cinza_escuro, cinza_claro, centro):
    from openpyxl.styles import Alignment

    aba = livro.create_sheet("Resumo Anual")
    aba.sheet_view.showGridLines = False
    aba.sheet_view.zoomScale = 90
    for coluna, largura in {"A": 9, "B": 39.4, "C": 33.4, "D": 23.3, "E": 23.3, "F": 23.3,
                            "G": 22.7, "H": 23.3, "I": 22.3}.items():
        aba.column_dimensions[coluna].width = largura
    tnr = "Times New Roman"
    moeda = '"R$"\\ #,##0.00'
    esquerda = Alignment(horizontal="left", vertical="center")

    def cab_meses(linha, col_inicial, meses):
        for j, mes in enumerate(meses):
            col = chr(ord(col_inicial) + j)
            escrever(aba, f"{col}{linha}", datetime(ano, MESES.index(mes) + 1, 1), tnr,
                     negrito=True, fundo=cinza_escuro, formato="mmm-yy", alinhamento=centro)

    def bloco(linha, titulo_c, titulo_meses, meses, linhas, primeira_col):
        """Bloco de títulos + colunas de meses (receita ou despesa)."""
        escrever(aba, f"B{linha}", "TÍTULOS", tnr, negrito=True, fundo=cinza_escuro, alinhamento=centro)
        escrever(aba, f"B{linha + 1}", None, tnr, fundo=cinza_escuro)
        aba.merge_cells(f"B{linha}:B{linha + 1}")
        if titulo_c:
            escrever(aba, f"C{linha}", titulo_c, tnr, negrito=True, fundo=cinza_escuro, alinhamento=centro)
            escrever(aba, f"C{linha + 1}", None, tnr, fundo=cinza_escuro)
            aba.merge_cells(f"C{linha}:C{linha + 1}")
        ultima = chr(ord(primeira_col) + len(meses) - 1)
        escrever(aba, f"{primeira_col}{linha}", titulo_meses, tnr, negrito=True, fundo=cinza_escuro, alinhamento=centro)
        for col in range(ord(primeira_col) + 1, ord(ultima) + 1):
            escrever(aba, f"{chr(col)}{linha}", None, tnr, fundo=cinza_escuro)
        aba.merge_cells(f"{primeira_col}{linha}:{ultima}{linha}")
        cab_meses(linha + 1, primeira_col, meses)
        for k, (titulo, destaque, formulas) in enumerate(linhas):
            r = linha + 2 + k
            fundo = cinza_claro if destaque else None
            escrever(aba, f"B{r}", titulo, tnr, negrito=destaque, fundo=fundo)
            for col, formula in formulas(r).items():
                escrever(aba, f"{col}{r}", formula, tnr, negrito=destaque, fundo=fundo, formato=moeda, alinhamento=centro)

    def por_mes(meses, primeira_col, celula):
        return {chr(ord(primeira_col) + j): f"={m}!{celula}" for j, m in enumerate(meses)}

    def soma(meses, primeira_col, expr):
        return {chr(ord(primeira_col) + j): expr(chr(ord(primeira_col) + j)) for j in range(len(meses))}

    titulo_receita = f"ARRECADAÇÃO CONFORME BALANCETES ANALISADOS NO PERÍODO DE JAN A DEZ DE {ano}"
    titulo_despesa = f"DESPESAS CONFORME BALANCETES ANALISADOS NO PERÍODO DE JAN A DEZ DE {ano}"
    s1, s2 = MESES[:6], MESES[6:]

    def receita(meses, primeira_col, previsao, acumulado_anterior):
        def acumulado(r):
            out = {}
            for j in range(len(meses)):
                col = chr(ord(primeira_col) + j)
                ant = (chr(ord(col) - 1) + str(r)) if j else acumulado_anterior
                out[col] = f"=SUM({col}{r - 4}:{col}{r - 1})" + (f"+{ant}" if ant else "")
            if previsao:
                out["C"] = f"=SUM(C{r - 4}:C{r - 1})"
            return out
        extra = lambda formulas: ({"C": formulas} if previsao else {})
        return [
            ("Receitas Correntes", False, lambda r: {**extra("=dez!D5"), **por_mes(meses, primeira_col, "E5")}),
            ("Total", True, lambda r: {**extra(f"=SUM(C{r - 1})"), **soma(meses, primeira_col, lambda c: f"=SUM({c}{r - 1})")}),
            ("Superávit Fin. Exer. Anterior", False, lambda r: {**extra("=dez!D7"), **por_mes(meses, primeira_col, "E7")}),
            ("Excesso de Arrecadação", False, lambda r: {**extra("=dez!D8"), **por_mes(meses, primeira_col, "E8")}),
            ("Saldo do Mês Anterior", False, lambda r: {**extra("=dez!D9"), **por_mes(meses, primeira_col, "E9")}),
            ("Total Acumulado", True, acumulado),
        ]

    bloco(3, "PREVISÃO ATUALIZADA", titulo_receita, s1, receita(s1, "D", True, None), "D")
    bloco(12, None, titulo_receita, s2, receita(s2, "C", False, "I10"), "C")

    def despesa(meses, primeira_col, dotacao):
        extra = lambda f: ({"C": f} if dotacao else {})
        return [
            ("Despesas Correntes", False, lambda r: {**extra("=dez!D16"), **por_mes(meses, primeira_col, "E16")}),
            ("Despesas de Capital", False, lambda r: {**extra("=dez!D17"), **por_mes(meses, primeira_col, "E17")}),
            ("Total", True, lambda r: {**extra(f"=SUM(C{r - 2}:C{r - 1})"),
                                       **soma(meses, primeira_col, lambda c: f"=SUM({c}{r - 2}:{c}{r - 1})")}),
            ("Suplementação", False, lambda r: {**extra("=dez!D20"), **por_mes(meses, primeira_col, "D20")}),
            ("Anulações", False, lambda r: {**extra("=dez!D22"), **por_mes(meses, primeira_col, "D22")}),
            ("Total Acumulado", True, lambda r: {**extra("=dez!D23"), **por_mes(meses, primeira_col, "E23")}),
        ]

    bloco(22, "DOTAÇÃO ATUALIZADA", titulo_despesa, s1, despesa(s1, "D", True), "D")
    bloco(31, None, titulo_despesa, s2, despesa(s2, "C", False), "C")

    def financeiro(linha, meses, titulo, total, itens):
        escrever(aba, f"B{linha}", titulo, tnr, negrito=True, fundo=cinza_escuro, alinhamento=centro)
        cab_meses(linha, "C", meses)
        escrever(aba, f"B{linha + 1}", total, tnr, negrito=True, fundo=cinza_claro, alinhamento=centro)
        for j in range(len(meses)):
            col = chr(ord("C") + j)
            escrever(aba, f"{col}{linha + 1}", f"=SUM({col}{linha + 2}:{col}{linha + 6})", tnr,
                     negrito=True, fundo=cinza_claro, formato=moeda, alinhamento=centro)
        for k, (rotulo, celula) in enumerate(itens):
            r = linha + 2 + k
            escrever(aba, f"B{r}", rotulo, tnr, alinhamento=esquerda)
            for col, formula in por_mes(meses, "C", celula).items():
                escrever(aba, f"{col}{r}", formula, tnr, formato=moeda, alinhamento=centro)

    ingressos = [("Receita Orçamentária", "K5"), ("Transf. Recebida Orçamentária", "K6"),
                 ("Receita Extraorçamentária", "K7"), ("Transf. Recebida Extraorçamentária", "K8"),
                 ("Saldo Anterior", "K9")]
    dispendios = [("Despesa Orçamentária", "M5"), ("Transf. Concedida Orçamentária", "M6"),
                  ("Despesa Extraorçamentária", "M7"), ("Transf. Concedida Extraorçamentária", "M8"),
                  ("Saldo para o Exercício Seguinte", "M9")]
    financeiro(41, s1, "RECEITA", "INGRESSOS", ingressos)
    financeiro(48, s1, "DESPESA", "DISPÊNDIOS", dispendios)
    financeiro(57, s2, "RECEITA", "INGRESSOS", ingressos)
    financeiro(64, s2, "DESPESA", "DISPÊNDIOS", dispendios)


def modelo_vazio() -> dict[str, dict]:
    return {mes: {campo: ZERO for campo in CAMPOS} for mes in MESES}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="base-calculo", description="Gera a base de cálculo mensal no modelo FEMTA.")
    parser.add_argument("--de", help="base existente (.xlsx) para copiar os valores digitados")
    parser.add_argument("--ano", type=int, help="exercício (padrão: o da base existente ou o ano atual)")
    parser.add_argument("--entidade", default="", help="nome do fundo para o título das abas")
    parser.add_argument("--saida", required=True, help="arquivo .xlsx a gerar")
    args = parser.parse_args(argv)

    try:
        base = ler_base(args.de) if args.de else {"ano": None, "meses": modelo_vazio()}
    except (ErroBase, FileNotFoundError) as erro:
        print(f"Erro ao ler a base: {erro}", file=sys.stderr)
        return 1
    ano = args.ano or base["ano"] or datetime.now().year
    gerar_base(base["meses"], args.saida, ano, args.entidade)
    print(f"Base de cálculo {ano} gerada: {args.saida}")
    for aviso in conferir(base["meses"]):
        print(f"  Aviso: {aviso}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
