"""Leitura e validação dos arquivos de entrada (CSV ou Excel)."""

from __future__ import annotations

import csv
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path


class ErroDeDados(Exception):
    """Erro nos dados de entrada. Guarda todos os problemas encontrados."""

    def __init__(self, problemas: list[str]):
        self.problemas = problemas
        super().__init__("\n".join(problemas))


@dataclass(frozen=True)
class Conta:
    codigo: str
    descricao: str
    natureza: str | None = None  # "D" (devedora), "C" (credora) ou None


@dataclass(frozen=True)
class Lancamento:
    data: date
    debito: str
    credito: str
    valor: Decimal
    historico: str = ""
    linha: int = 0


def _normalizar_cabecalho(nome: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode()
    return sem_acento.strip().lower().replace(" ", "_")


def ler_tabela(caminho: str | Path) -> list[dict[str, object]]:
    """Lê um CSV (separador ; ou ,) ou a primeira aba de um .xlsx.

    Devolve uma lista de dicionários com os cabeçalhos normalizados
    (minúsculas, sem acentos). Linhas totalmente vazias são ignoradas.
    """
    caminho = Path(caminho)
    if not caminho.exists():
        raise ErroDeDados([f"Arquivo não encontrado: {caminho}"])

    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        planilha = load_workbook(caminho, read_only=True, data_only=True).worksheets[0]
        linhas = list(planilha.iter_rows(values_only=True))
    else:
        with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
            amostra = arquivo.read(4096)
            arquivo.seek(0)
            separador = ";" if amostra.count(";") >= amostra.count(",") else ","
            linhas = list(csv.reader(arquivo, delimiter=separador))

    if not linhas:
        return []

    cabecalho = [_normalizar_cabecalho(c) if c is not None else "" for c in linhas[0]]
    registros = []
    for valores in linhas[1:]:
        if all(v is None or str(v).strip() == "" for v in valores):
            continue
        registros.append(dict(zip(cabecalho, valores)))
    return registros


def converter_valor(valor: object) -> Decimal:
    """Converte "1.234,56", "1234.56", "R$ 10,00" ou números em Decimal."""
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        return Decimal(str(valor))
    texto = str(valor or "").replace("R$", "").replace(" ", "").strip()
    if not texto:
        raise ValueError("valor vazio")
    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()")
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        numero = Decimal(texto)
    except InvalidOperation:
        raise ValueError(f"valor inválido: {valor!r}") from None
    return -numero if negativo else numero


def converter_data(valor: object) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ValueError(f"data inválida: {valor!r}")


def _texto(registro: dict, campo: str) -> str:
    valor = registro.get(campo)
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _exigir_colunas(registros: list[dict], colunas: list[str], arquivo: str) -> None:
    if registros:
        faltando = [c for c in colunas if c not in registros[0]]
        if faltando:
            raise ErroDeDados([f"{arquivo}: coluna(s) obrigatória(s) ausente(s): {', '.join(faltando)}"])


def ler_plano_de_contas(caminho: str | Path) -> dict[str, Conta]:
    """Colunas: codigo, descricao e (opcional) natureza (D/C)."""
    registros = ler_tabela(caminho)
    _exigir_colunas(registros, ["codigo", "descricao"], "Plano de contas")

    contas: dict[str, Conta] = {}
    problemas = []
    for numero, registro in enumerate(registros, start=2):
        codigo = _texto(registro, "codigo")
        natureza = _texto(registro, "natureza").upper()[:1] or None
        if not codigo:
            problemas.append(f"Plano de contas, linha {numero}: código vazio")
            continue
        if codigo in contas:
            problemas.append(f"Plano de contas, linha {numero}: código {codigo} repetido")
            continue
        if natureza not in (None, "D", "C"):
            problemas.append(f"Plano de contas, linha {numero}: natureza deve ser D ou C")
            continue
        contas[codigo] = Conta(codigo, _texto(registro, "descricao"), natureza)

    if problemas:
        raise ErroDeDados(problemas)
    if not contas:
        raise ErroDeDados(["Plano de contas vazio"])
    return contas


def ler_lancamentos(caminho: str | Path) -> list[Lancamento]:
    """Colunas: data, debito, credito, valor e (opcional) historico."""
    registros = ler_tabela(caminho)
    _exigir_colunas(registros, ["data", "debito", "credito", "valor"], "Lançamentos")

    lancamentos = []
    problemas = []
    for numero, registro in enumerate(registros, start=2):
        try:
            lancamentos.append(
                Lancamento(
                    data=converter_data(registro.get("data")),
                    debito=_texto(registro, "debito"),
                    credito=_texto(registro, "credito"),
                    valor=converter_valor(registro.get("valor")),
                    historico=_texto(registro, "historico"),
                    linha=numero,
                )
            )
        except ValueError as erro:
            problemas.append(f"Lançamentos, linha {numero}: {erro}")

    if problemas:
        raise ErroDeDados(problemas)
    return lancamentos


def ler_saldos_iniciais(caminho: str | Path) -> dict[str, Decimal]:
    """Colunas: codigo, saldo e (opcional) dc (D/C).

    Devolve saldos com sinal: positivo = devedor, negativo = credor.
    Sem a coluna dc, o próprio sinal do saldo é usado.
    """
    registros = ler_tabela(caminho)
    _exigir_colunas(registros, ["codigo", "saldo"], "Saldos iniciais")

    saldos: dict[str, Decimal] = {}
    problemas = []
    for numero, registro in enumerate(registros, start=2):
        codigo = _texto(registro, "codigo")
        try:
            saldo = converter_valor(registro.get("saldo"))
        except ValueError as erro:
            problemas.append(f"Saldos iniciais, linha {numero}: {erro}")
            continue
        dc = _texto(registro, "dc").upper()[:1]
        if dc == "C":
            saldo = -abs(saldo)
        elif dc == "D":
            saldo = abs(saldo)
        saldos[codigo] = saldos.get(codigo, Decimal(0)) + saldo

    if problemas:
        raise ErroDeDados(problemas)
    return saldos
