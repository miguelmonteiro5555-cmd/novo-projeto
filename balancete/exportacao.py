"""Exportação do balancete para Excel (.xlsx) e CSV."""

from __future__ import annotations

import csv
from decimal import Decimal
from pathlib import Path

from .calculo import Balancete
from .formato import formatar_valor

MESES = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]

COLUNAS = [
    "Código", "Descrição", "Saldo anterior", "D/C",
    "Débitos", "Créditos", "Saldo atual", "D/C",
]


def _dc(valor: Decimal) -> str:
    return "" if valor == 0 else ("D" if valor > 0 else "C")


def _valores(linha) -> list:
    return [
        linha.conta.codigo,
        linha.conta.descricao,
        abs(linha.saldo_anterior), _dc(linha.saldo_anterior),
        linha.debitos,
        linha.creditos,
        abs(linha.saldo_atual), _dc(linha.saldo_atual),
    ]


def titulo(balancete: Balancete, empresa: str = "") -> str:
    texto = f"Balancete de Verificação — {MESES[balancete.mes - 1]}/{balancete.ano}"
    return f"{empresa} — {texto}" if empresa else texto


def exportar_excel(balancete: Balancete, caminho: str | Path, empresa: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro = Workbook()
    aba = livro.active
    aba.title = f"Balancete {balancete.mes:02d}-{balancete.ano}"

    formato_numero = '#,##0.00;-#,##0.00;"-"'
    negrito = Font(bold=True)
    fundo_cabecalho = PatternFill("solid", fgColor="1F4E78")
    fundo_sintetica = PatternFill("solid", fgColor="DDEBF7")
    borda_total = Border(top=Side(style="thin"), bottom=Side(style="double"))

    aba["A1"] = titulo(balancete, empresa)
    aba["A1"].font = Font(bold=True, size=14)
    aba.merge_cells("A1:H1")

    aba.append([])
    aba.append(COLUNAS)
    for celula in aba[3]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = fundo_cabecalho
        celula.alignment = Alignment(horizontal="center")

    for linha in balancete.linhas:
        aba.append(_valores(linha))
        atual = aba.max_row
        aba.cell(atual, 2).alignment = Alignment(indent=linha.nivel - 1)
        for coluna in (3, 5, 6, 7):
            aba.cell(atual, coluna).number_format = formato_numero
        if linha.sintetica:
            for celula in aba[atual]:
                celula.font = negrito
                celula.fill = fundo_sintetica

    aba.append([
        "", "TOTAL",
        abs(balancete.total_saldo_anterior), _dc(balancete.total_saldo_anterior),
        balancete.total_debitos, balancete.total_creditos,
        abs(balancete.total_saldo_atual), _dc(balancete.total_saldo_atual),
    ])
    total = aba.max_row
    for celula in aba[total]:
        celula.font = negrito
        celula.border = borda_total
    for coluna in (3, 5, 6, 7):
        aba.cell(total, coluna).number_format = formato_numero

    aba.append([])
    situacao = (
        "✔ Balancete fechado: débitos = créditos e saldos somam zero."
        if balancete.fechado
        else "✘ Balancete NÃO fechou — veja a aba Avisos."
    )
    aba.append([situacao])
    aba.cell(aba.max_row, 1).font = Font(bold=True, color="2E7D32" if balancete.fechado else "C62828")

    for coluna, largura in zip("ABCDEFGH", (14, 45, 16, 5, 16, 16, 16, 5)):
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = "A4"

    if balancete.avisos:
        avisos = livro.create_sheet("Avisos")
        avisos.append(["Avisos"])
        avisos["A1"].font = negrito
        for aviso in balancete.avisos:
            avisos.append([aviso])
        avisos.column_dimensions["A"].width = 110

    livro.save(caminho)
    return caminho


def exportar_csv(balancete: Balancete, caminho: str | Path) -> Path:
    """CSV com separador ; e vírgula decimal (abre direto no Excel em pt-BR)."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding="utf-8-sig", newline="") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(COLUNAS)
        for linha in balancete.linhas:
            escritor.writerow(
                formatar_valor(v) if isinstance(v, Decimal) else v for v in _valores(linha)
            )
    return caminho
