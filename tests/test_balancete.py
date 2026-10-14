from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from balancete import (
    Conta,
    ErroDeDados,
    Lancamento,
    exportar_csv,
    exportar_excel,
    gerar_balancete,
    ler_lancamentos,
    ler_plano_de_contas,
    ler_saldos_iniciais,
)
from balancete.cli import main
from balancete.leitura import converter_valor

EXEMPLOS = Path(__file__).resolve().parent.parent / "exemplos"

PLANO = {
    "1": Conta("1", "ATIVO", "D"),
    "1.1": Conta("1.1", "Caixa e Bancos", "D"),
    "1.1.01": Conta("1.1.01", "Caixa", "D"),
    "1.1.02": Conta("1.1.02", "Banco", "D"),
    "2": Conta("2", "PASSIVO", "C"),
    "2.1": Conta("2.1", "Capital Social", "C"),
    "3": Conta("3", "RECEITAS", "C"),
    "3.1": Conta("3.1", "Vendas", "C"),
}


def lanc(dia: date, debito: str, credito: str, valor: str) -> Lancamento:
    return Lancamento(dia, debito, credito, Decimal(valor))


def linha(balancete, codigo):
    return next(l for l in balancete.linhas if l.conta.codigo == codigo)


@pytest.mark.parametrize(
    "texto, esperado",
    [
        ("1.234,56", "1234.56"),
        ("1234.56", "1234.56"),
        ("1,234.56", "1234.56"),
        ("R$ 10,00", "10.00"),
        ("(50,00)", "-50.00"),
        (12.5, "12.5"),
    ],
)
def test_converter_valor(texto, esperado):
    assert converter_valor(texto) == Decimal(esperado)


def test_saldo_anterior_movimento_e_saldo_atual():
    lancamentos = [
        lanc(date(2026, 8, 10), "1.1.01", "2.1", "1000"),  # mês anterior
        lanc(date(2026, 9, 5), "1.1.02", "3.1", "300"),
        lanc(date(2026, 9, 6), "1.1.01", "1.1.02", "100"),
        lanc(date(2026, 10, 1), "1.1.01", "3.1", "999"),  # mês seguinte: ignorado
    ]
    b = gerar_balancete(PLANO, lancamentos, 2026, 9)

    caixa = linha(b, "1.1.01")
    assert (caixa.saldo_anterior, caixa.debitos, caixa.creditos, caixa.saldo_atual) == (1000, 100, 0, 1100)
    banco = linha(b, "1.1.02")
    assert (banco.saldo_anterior, banco.debitos, banco.creditos, banco.saldo_atual) == (0, 300, 100, 200)
    vendas = linha(b, "3.1")
    assert vendas.saldo_atual == -300  # credor

    assert b.total_debitos == b.total_creditos == 400
    assert b.fechado


def test_contas_sinteticas_somam_as_filhas():
    lancamentos = [
        lanc(date(2026, 9, 1), "1.1.01", "2.1", "500"),
        lanc(date(2026, 9, 2), "1.1.02", "2.1", "250"),
    ]
    b = gerar_balancete(PLANO, lancamentos, 2026, 9)

    for codigo in ("1", "1.1"):
        assert linha(b, codigo).sintetica
        assert linha(b, codigo).debitos == 750
    assert linha(b, "1.1").nivel == 2
    assert not linha(b, "1.1.01").sintetica
    # sintéticas não entram duas vezes no total
    assert b.total_debitos == 750


def test_saldos_iniciais_e_contas_zeradas():
    saldos = {"1.1.01": Decimal("100"), "2.1": Decimal("-100")}
    b = gerar_balancete(PLANO, [], 2026, 9, saldos)
    codigos = [l.conta.codigo for l in b.linhas]
    assert "3.1" not in codigos  # zerada é omitida por padrão
    assert linha(b, "1.1.01").saldo_atual == 100
    assert b.fechado

    completo = gerar_balancete(PLANO, [], 2026, 9, saldos, incluir_zeradas=True)
    assert len(completo.linhas) == len(PLANO)


def test_ordenacao_numerica_dos_codigos():
    plano = {c: Conta(c, c) for c in ["1", "1.10", "1.2", "1.9"]}
    b = gerar_balancete(plano, [lanc(date(2026, 1, 1), "1.10", "1.2", "1")], 2026, 1, incluir_zeradas=True)
    assert [l.conta.codigo for l in b.linhas] == ["1", "1.2", "1.9", "1.10"]


def test_erros_de_validacao_sao_todos_reportados():
    lancamentos = [
        Lancamento(date(2026, 9, 1), "9.9", "2.1", Decimal("10"), linha=2),
        Lancamento(date(2026, 9, 1), "1.1", "2.1", Decimal("10"), linha=3),
        Lancamento(date(2026, 9, 1), "1.1.01", "1.1.01", Decimal("10"), linha=4),
        Lancamento(date(2026, 9, 1), "1.1.01", "2.1", Decimal("-5"), linha=5),
    ]
    with pytest.raises(ErroDeDados) as erro:
        gerar_balancete(PLANO, lancamentos, 2026, 9)
    texto = str(erro.value)
    assert "linha 2 (débito): conta 9.9 não existe" in texto
    assert "linha 3 (débito): conta 1.1 é sintética" in texto
    assert "linha 4: débito e crédito na mesma conta" in texto
    assert "linha 5: valor deve ser maior que zero" in texto


def test_avisos_saldo_invertido_e_saldos_iniciais_desbalanceados():
    b = gerar_balancete(PLANO, [], 2026, 9, {"1.1.01": Decimal("-50")})
    assert not b.fechado
    assert any("saldos anteriores não é zero" in a for a in b.avisos)
    assert any("1.1.01" in a and "saldo invertido" in a for a in b.avisos)


def test_leitura_dos_exemplos_csv():
    plano = ler_plano_de_contas(EXEMPLOS / "plano_de_contas.csv")
    lancamentos = ler_lancamentos(EXEMPLOS / "lancamentos.csv")
    saldos = ler_saldos_iniciais(EXEMPLOS / "saldos_iniciais.csv")

    assert plano["1.1.01"] == Conta("1.1.01", "Caixa", "D")
    assert lancamentos[0].data == date(2026, 8, 5)
    assert lancamentos[0].valor == Decimal("8000.00")
    assert saldos["2.3.01"] == Decimal("-90000.00")
    assert sum(saldos.values()) == 0


def test_leitura_de_xlsx_com_cabecalho_acentuado(tmp_path):
    livro = Workbook()
    aba = livro.active
    aba.append(["Data", "Débito", "Crédito", "Valor", "Histórico"])
    aba.append([date(2026, 9, 3), "1.1.01", "3.1", 150.75, "Venda"])
    caminho = tmp_path / "lancamentos.xlsx"
    livro.save(caminho)

    (lancamento,) = ler_lancamentos(caminho)
    assert lancamento.debito == "1.1.01"
    assert lancamento.credito == "3.1"
    assert lancamento.valor == Decimal("150.75")
    assert lancamento.data == date(2026, 9, 3)


def test_coluna_obrigatoria_ausente(tmp_path):
    caminho = tmp_path / "lanc.csv"
    caminho.write_text("data;debito;valor\n01/09/2026;1.1.01;10\n", encoding="utf-8")
    with pytest.raises(ErroDeDados, match="credito"):
        ler_lancamentos(caminho)


def test_exportacoes(tmp_path):
    b = gerar_balancete(PLANO, [lanc(date(2026, 9, 1), "1.1.01", "3.1", "1234.5")], 2026, 9)

    xlsx = exportar_excel(b, tmp_path / "b.xlsx", empresa="ACME")
    aba = load_workbook(xlsx).active
    assert aba["A1"].value == "ACME — Balancete de Verificação — Setembro/2026"
    valores = [r for r in aba.iter_rows(values_only=True) if r[1] == "TOTAL"]
    assert valores[0][4] == valores[0][5] == 1234.5

    texto = exportar_csv(b, tmp_path / "b.csv").read_text(encoding="utf-8-sig")
    assert "1.1.01;Caixa;0,00;;1.234,50;0,00;1.234,50;D" in texto


def test_cli_com_exemplos(tmp_path, capsys):
    saida = tmp_path / "balancete.xlsx"
    codigo = main([
        "--mes", "09/2026",
        "--plano", str(EXEMPLOS / "plano_de_contas.csv"),
        "--lancamentos", str(EXEMPLOS / "lancamentos.csv"),
        "--saldos", str(EXEMPLOS / "saldos_iniciais.csv"),
        "--saida", str(saida),
    ])
    assert codigo == 0
    assert saida.exists()
    assert "FECHADO" in capsys.readouterr().out


def test_cli_reporta_erros(tmp_path, capsys):
    lancamentos = tmp_path / "lanc.csv"
    lancamentos.write_text("data;debito;credito;valor\n01/09/2026;9.9;1.1.01;10\n", encoding="utf-8")
    codigo = main([
        "--mes", "2026-09",
        "--plano", str(EXEMPLOS / "plano_de_contas.csv"),
        "--lancamentos", str(lancamentos),
        "--saida", str(tmp_path / "x.xlsx"),
    ])
    assert codigo == 1
    assert "conta 9.9 não existe" in capsys.readouterr().err
