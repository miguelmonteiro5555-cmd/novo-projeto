import shutil
import subprocess
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from balancete.base_calculo import (
    MESES,
    _numero,
    calcular,
    conferir,
    gerar_base,
    ler_base,
    main,
    modelo_vazio,
)


def dados_exemplo():
    """Dois meses com balanço financeiro fechado e dotação conferida."""
    meses = modelo_vazio()
    jan, fev = meses["jan"], meses["fev"]
    for mes in (jan, fev):
        mes.update(receita_prevista=Decimal("1000"), correntes_inicial=Decimal("700"), capital_inicial=Decimal("300"))
    jan.update(
        receita_atualizada=Decimal("1000"), receita_realizada=Decimal("200"),
        correntes_autorizada=Decimal("700"), capital_autorizada=Decimal("300"),
        correntes_execucao=Decimal("120"), capital_execucao=Decimal("30"),
        receita_extraorcamentaria=Decimal("50"), despesa_extraorcamentaria=Decimal("40"),
        saldo_exercicio_anterior=Decimal("100"),
        saldo_final=Decimal("160"),  # 100 + 200 + 50 - 150 - 40
    )
    fev.update(
        receita_atualizada=Decimal("1100"), receita_realizada=Decimal("300"),
        correntes_autorizada=Decimal("760"), capital_autorizada=Decimal("290"),
        suplementacao_autorizada=Decimal("80"), anulacoes_autorizada=Decimal("30"),
        correntes_execucao=Decimal("200"), capital_execucao=Decimal("50"), anulacoes_execucao=Decimal("10"),
        receita_extraorcamentaria=Decimal("20"), despesa_extraorcamentaria=Decimal("30"),
        saldo_final=Decimal("210"),  # 160 + 300 + 20 - 240 - 30
    )
    return meses


@pytest.mark.parametrize("valor, esperado", [
    (None, "0"), (12.5, "12.5"), ("=3745677.7+66760.56", "3812438.26"),
    ("=10-2.5+1", "8.5"), ("1.234,56", "1234.56"), ("  ", "0"),
])
def test_numero(valor, esperado):
    assert _numero(valor) == Decimal(esperado)


def test_numero_rejeita_formula_com_referencia():
    with pytest.raises(ValueError):
        _numero("=A1+2")


def test_calculo_do_mes_e_acumulados():
    r = calcular(dados_exemplo())
    jan, fev = r["jan"], r["fev"]
    assert jan["ingressos"] == jan["dispendios"] == Decimal("350")
    assert jan["rfp"] == Decimal("60")  # (200 + 50) - (150 + 40)
    assert fev["saldo_anterior"] == Decimal("160")  # saldo final de janeiro
    assert fev["despesa_orcamentaria"] == Decimal("240")
    assert fev["conferencia_balanco"] == 0
    assert fev["conferencia_dotacao"] == 0
    assert fev["receita_acumulada"] == Decimal("500")
    assert fev["despesa_acumulada"] == Decimal("390")
    assert fev["diferenca_quadros"] == Decimal("50")


def test_conferencias_apontam_diferencas():
    meses = dados_exemplo()
    assert conferir(meses) == []
    meses["fev"]["saldo_final"] = Decimal("200")
    meses["fev"]["anulacoes_autorizada"] = Decimal("0")
    avisos = conferir(meses)
    assert any(a.startswith("fev: ingressos e dispêndios diferem em 10.00") for a in avisos)
    assert any(a.startswith("fev: dotação autorizada difere") for a in avisos)


def test_gerar_e_ler_de_volta(tmp_path):
    meses = dados_exemplo()
    caminho = gerar_base(meses, tmp_path / "base.xlsx", 2026, "Fundo Teste")

    livro = load_workbook(caminho)
    assert livro.sheetnames == MESES + ["Resumo Anual"]
    fev = livro["fev"]
    assert fev["K9"].value == "=jan!M9"
    assert fev["J9"].value == "Saldo do Mês Anterior"
    assert livro["jan"]["J9"].value == "Saldo do Exercício Anterior"
    assert fev["E23"].value == "=jan!E23+fev!E18-fev!E22"
    assert fev["D23"].value == "=D18+D19"
    assert livro["Resumo Anual"]["D5"].value == "=jan!E5"
    assert "2026" in livro["Resumo Anual"]["D3"].value

    lido = ler_base(caminho)
    assert lido["ano"] == 2026
    for mes in ("jan", "fev"):
        for campo, valor in meses[mes].items():
            if mes == "fev" and campo == "saldo_exercicio_anterior":
                continue
            assert lido["meses"][mes][campo] == valor, (mes, campo)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice não instalado")
def test_formulas_batem_com_o_calculo(tmp_path):
    meses = dados_exemplo()
    gerar_base(meses, tmp_path / "base.xlsx", 2026)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                    str(tmp_path / "recalc"), str(tmp_path / "base.xlsx")], check=True, capture_output=True, timeout=120)
    livro = load_workbook(tmp_path / "recalc" / "base.xlsx", data_only=True)
    esperado = calcular(meses)
    for mes in ("jan", "fev"):
        aba, r = livro[mes], esperado[mes]
        assert aba["K4"].value == pytest.approx(float(r["ingressos"]))
        assert aba["M4"].value == pytest.approx(float(r["dispendios"]))
        assert aba["J17"].value == pytest.approx(float(r["rfp"]))
        assert aba["J21"].value == pytest.approx(float(r["diferenca_quadros"]))
        assert aba["E10"].value == pytest.approx(float(r["receita_acumulada"]))
        assert aba["E23"].value == pytest.approx(float(r["despesa_acumulada"]))
        assert aba["M17"].value == "OK"
    assert livro["Resumo Anual"]["D42"].value == pytest.approx(float(esperado["fev"]["ingressos"]))


def test_cli_modelo_vazio(tmp_path, capsys):
    saida = tmp_path / "modelo.xlsx"
    assert main(["--ano", "2027", "--entidade", "Fundo X", "--saida", str(saida)]) == 0
    assert "2027" in capsys.readouterr().out
    assert load_workbook(saida)["jan"]["B1"].value == "Fundo X — Base de cálculo JAN/2027"
